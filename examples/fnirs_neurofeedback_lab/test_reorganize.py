#!/usr/bin/env python3
"""
Hyperscanning Toolkit

Copyright (c) 2026 Dr. Yael Hodaya Moshe.

Lead Developer:
    Dr. Yael Hodaya Moshe

Developed in collaboration with the Social Neuroscience Lab.

Scientific Supervision:
    Dr. Hila Gvirts
    Dr. Anat Dahan

This file is part of the Hyperscanning Toolkit.

---

TEST VERSION: Reorganize a small subset of files to verify structure.
This processes only the first N files from each condition as a proof-of-concept.

Usage:
    python test_reorganize.py \\
        --nfb-dir "<path to NFB O2HB Data>" \\
        --control-dir "<path to Control O2HB Data>" \\
        [--output-dir <path, default: ../../data_test next to this script>] \\
        [--limit 10]
"""

import argparse
import csv
import shutil
from collections import defaultdict
from pathlib import Path

import openpyxl

DEFAULT_OUTPUT_DIR = Path(__file__).resolve().parent / ".." / ".." / "data_test"

# Configuration, set from CLI args in main()
SOURCE_DIRS = {}
OUTPUT_BASE = None

# Data structures
summary_data = []
dyad_manifest = []
warnings_list = []


def log_warning(msg):
    """Log a warning message"""
    warnings_list.append(msg)
    print(f"  ⚠ WARNING: {msg}")


def extract_metadata(filepath):
    """Extract Group, Session, Subject from Excel file"""
    try:
        wb = openpyxl.load_workbook(filepath, data_only=True)
        ws = wb.active

        headers = [cell.value for cell in ws[1]]

        if 'Group' not in headers or 'Session' not in headers or 'Subject' not in headers:
            log_warning(f"{filepath.name} - Missing expected headers")
            wb.close()
            return None

        data_row = [cell.value for cell in ws[2]]

        try:
            group = int(data_row[headers.index('Group')])
            session = int(data_row[headers.index('Session')])
            subject = int(data_row[headers.index('Subject')])
        except (ValueError, TypeError, IndexError) as e:
            log_warning(f"{filepath.name} - Error parsing metadata: {e}")
            wb.close()
            return None

        wb.close()
        return {'group': group, 'session': session, 'subject': subject}

    except Exception as e:
        log_warning(f"{filepath.name} - Failed to read: {e}")
        return None


def reorganize_files(condition, source_dir, limit=10):
    """Reorganize files for a given condition (NFB or Control)"""
    print(f"\n{'='*80}")
    print(f"Processing {condition} files (limit: {limit})...")
    print(f"{'='*80}")

    if not source_dir.exists():
        print(f"ERROR: Source directory not found: {source_dir}")
        return

    excel_files = sorted(source_dir.glob("*.xlsx"))[:limit]
    print(f"Processing {len(excel_files)} files from {condition}\n")

    dyad_files = defaultdict(lambda: defaultdict(dict))

    # Process each file
    for fpath in excel_files:
        metadata = extract_metadata(fpath)

        if metadata is None:
            continue

        group = metadata['group']
        session = metadata['session']
        subject = metadata['subject']

        if subject not in [1, 2]:
            log_warning(f"{fpath.name} - Invalid subject value: {subject}")
            continue

        dyad_files[group][session][subject] = {
            'filename': fpath.name,
            'path': fpath,
            'subject': subject
        }

        print(f"  ✓ {fpath.name}")
        print(f"    Group {group}, Session {session}, Subject {subject}")

    # Create output directories and copy files
    print(f"\n{'='*80}")
    print(f"Creating directories and copying files...")
    print(f"{'='*80}\n")

    output_condition_dir = OUTPUT_BASE / condition

    for group in sorted(dyad_files.keys()):
        dyad_id = f"{group:04d}"
        dyad_dir = output_condition_dir / f"dyad_{dyad_id}"

        for session in sorted(dyad_files[group].keys()):
            session_id = f"{session:02d}"
            session_dir = dyad_dir / f"session_{session_id}"
            session_dir.mkdir(parents=True, exist_ok=True)

            subjects_in_session = dyad_files[group][session]

            if 1 not in subjects_in_session and 2 not in subjects_in_session:
                log_warning(f"Group {group} Session {session} - No subjects found")
                continue

            if len(subjects_in_session) < 2:
                missing_subject = 2 if 1 in subjects_in_session else 1
                log_warning(f"Group {group} Session {session} - Missing subject {missing_subject}")

            for subject, file_info in subjects_in_session.items():
                src_path = file_info['path']
                dst_filename = f"subject_{subject:02d}.xlsx"
                dst_path = session_dir / dst_filename

                shutil.copy2(src_path, dst_path)
                print(f"  ✓ {condition}/dyad_{dyad_id}/session_{session_id}/{dst_filename}")

                summary_data.append({
                    'condition': condition,
                    'dyad_id': dyad_id,
                    'session': session_id,
                    'subject': subject,
                    'original_filename': src_path.name,
                    'new_filepath': str(dst_path.relative_to(OUTPUT_BASE))
                })

            subject_1_file = f"subject_01.xlsx" if 1 in subjects_in_session else "MISSING"
            subject_2_file = f"subject_02.xlsx" if 2 in subjects_in_session else "MISSING"

            dyad_manifest.append({
                'condition': condition,
                'dyad_id': dyad_id,
                'session': session_id,
                'subject1_file': subject_1_file,
                'subject2_file': subject_2_file
            })


def generate_reports():
    """Generate summary and manifest CSV files"""
    print(f"\n{'='*80}")
    print("Generating CSV reports...")
    print(f"{'='*80}\n")

    summary_path = OUTPUT_BASE / "summary.csv"
    with open(summary_path, 'w', newline='', encoding='utf-8') as f:
        if summary_data:
            fieldnames = list(summary_data[0].keys())
            writer = csv.DictWriter(f, fieldnames=fieldnames)
            writer.writeheader()
            writer.writerows(summary_data)
    print(f"  ✓ {summary_path.name} ({len(summary_data)} records)")

    manifest_path = OUTPUT_BASE / "dyad_manifest.csv"
    with open(manifest_path, 'w', newline='', encoding='utf-8') as f:
        if dyad_manifest:
            fieldnames = list(dyad_manifest[0].keys())
            writer = csv.DictWriter(f, fieldnames=fieldnames)
            writer.writeheader()
            writer.writerows(dyad_manifest)
    print(f"  ✓ {manifest_path.name} ({len(dyad_manifest)} records)")


def print_summary():
    """Print summary statistics"""
    print(f"\n{'='*80}")
    print("TEST SUMMARY")
    print(f"{'='*80}\n")

    print(f"Files processed: {len(summary_data)}")
    print(f"Summary records: {len(summary_data)}")
    print(f"Dyad-session combinations: {len(dyad_manifest)}")
    print(f"Warnings: {len(warnings_list)}")

    if warnings_list:
        print(f"\nWarnings detected:")
        for i, warning in enumerate(warnings_list, 1):
            print(f"  {i}. {warning}")

    print(f"\n\nOutput directory: {OUTPUT_BASE}")
    print(f"  {OUTPUT_BASE}/summary.csv")
    print(f"  {OUTPUT_BASE}/dyad_manifest.csv")


def parse_args():
    parser = argparse.ArgumentParser(
        description="TEST: reorganize a small subset of raw fNIRS Excel exports as a proof-of-concept."
    )
    parser.add_argument("--nfb-dir", required=True, help="Path to the raw 'NFB O2HB Data' folder")
    parser.add_argument("--control-dir", required=True, help="Path to the raw 'Control O2HB Data' folder")
    parser.add_argument(
        "--output-dir",
        default=str(DEFAULT_OUTPUT_DIR),
        help=f"Where to write the reorganized subset (default: {DEFAULT_OUTPUT_DIR})",
    )
    parser.add_argument("--limit", type=int, default=10, help="Number of files per condition to process (default: 10)")
    return parser.parse_args()


def main():
    """Main execution"""
    global SOURCE_DIRS, OUTPUT_BASE

    args = parse_args()
    SOURCE_DIRS = {'NFB': Path(args.nfb_dir), 'Control': Path(args.control_dir)}
    OUTPUT_BASE = Path(args.output_dir)

    print(f"\n{'='*80}")
    print("FNIRS DATASET REORGANIZER - TEST MODE")
    print(f"{'='*80}\n")

    OUTPUT_BASE.mkdir(parents=True, exist_ok=True)

    for condition, source_dir in SOURCE_DIRS.items():
        reorganize_files(condition, source_dir, limit=args.limit)

    generate_reports()
    print_summary()

    print(f"\n{'='*80}")
    print("TEST COMPLETE")
    print(f"{'='*80}\n")


if __name__ == "__main__":
    main()
